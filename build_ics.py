#!/usr/bin/env python3
"""
Download a SharePoint/OneDrive Excel timetable and build an .ics feed
for selected sections (default: S2, S4).

Env vars:
  SHEET_URL      SharePoint share link (required, unless --file is given)
  SECTIONS       comma-separated, default "S2,S4"
  TZ_NAME        timezone of the timetable, default "Asia/Kolkata"
  CAL_NAME       calendar display name
  DEFAULT_MINS   class length if no end time is found, default 90
  OUTPUT         output path, default "docs/calendar.ics"

Usage:
  python build_ics.py                 # normal run (used by GitHub Actions)
  python build_ics.py --file x.xlsx   # run on a local copy
  python build_ics.py --file x.xlsx --inspect   # print sheet layout to debug
"""
import argparse, datetime as dt, hashlib, io, os, re, sys
from zoneinfo import ZoneInfo

import openpyxl
import requests
from dateutil import parser as dparser

SECTIONS = [s.strip().upper() for s in os.environ.get("SECTIONS", "S2,S4").split(",") if s.strip()]
TZ = ZoneInfo(os.environ.get("TZ_NAME", "Asia/Kolkata"))
CAL_NAME = os.environ.get("CAL_NAME", f"MICA Classes ({', '.join(SECTIONS)})")
DEFAULT_MINS = int(os.environ.get("DEFAULT_MINS", "90"))
OUTPUT = os.environ.get("OUTPUT", "docs/calendar.ics")
ALL_WORDS = re.compile(r"\b(all\s+sections?|all\s+students|combined|common)\b", re.I)

# ---------------------------------------------------------------- download
def download(url: str) -> bytes:
    base = url.split("?")[0]
    tries = [base + "?download=1", url + ("&" if "?" in url else "?") + "download=1"]
    headers = {"User-Agent": "Mozilla/5.0 (ics-feed-bot)"}
    for u in tries:
        r = requests.get(u, headers=headers, allow_redirects=True, timeout=60)
        if r.ok and r.content[:2] == b"PK":          # xlsx files are zip archives
            return r.content
    sys.exit("ERROR: could not download the Excel file. The link is probably restricted "
             "to MICA accounts. Ask for an 'Anyone with the link' share link, or see README.")

# ---------------------------------------------------------------- helpers
def section_matches(text: str) -> list[str]:
    """Return which of SECTIONS appear in text (handles S2, S-2, Sec 2, S1-S4, 'all sections')."""
    if not text:
        return []
    t = str(text)
    if ALL_WORDS.search(t):
        return list(SECTIONS)
    found = set()
    for a, b in re.findall(r"(?<![A-Za-z0-9])S(?:ec(?:tion)?)?\s*[-.:]?\s*0?(\d{1,2})\s*(?:-|–|to)\s*S?(?:ec)?\s*0?(\d{1,2})(?!\d)", t, re.I):
        for n in range(int(a), int(b) + 1):
            found.add(f"S{n}")
    for n in re.findall(r"(?<![A-Za-z0-9])S(?:ec(?:tion)?)?\s*[-.:]?\s*0?(\d{1,2})(?!\d)", t, re.I):
        found.add(f"S{int(n)}")
    return [s for s in SECTIONS if s in found]

def to_date(v, year_hint):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 60000:       # Excel serial date
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    s = str(v).strip()
    if not re.search(r"\d", s):
        return None
    try:
        d = dparser.parse(s, dayfirst=True, fuzzy=True, default=dt.datetime(year_hint, 1, 1)).date()
        return d if 2000 < d.year < 2100 else None
    except (ValueError, OverflowError):
        return None

TIME_RE = re.compile(r"(\d{1,2})(?:[:.](\d{2}))?\s*([ap]\.?\s*m\.?)?", re.I)

def _fix(h, m, ap):
    if ap:
        pm = ap.lower().startswith("p")
        h = h % 12 + (12 if pm else 0)
    elif h < 8:                      # 1:00 in a class timetable means 13:00
        h += 12
    return dt.time(h % 24, m)

def parse_times(v):
    """Return (start, end|None) from a cell like '9:00-10:30', '2.15 PM to 3.45 PM', or a time object."""
    if v is None:
        return None, None
    if isinstance(v, dt.datetime):
        return v.time(), None
    if isinstance(v, dt.time):
        return v, None
    if isinstance(v, float) and 0 <= v < 1:
        mins = round(v * 1440)
        return dt.time(mins // 60, mins % 60), None
    s = str(v)
    hits = [m for m in TIME_RE.finditer(s) if m.group(2) or m.group(3)]
    if not hits:
        return None, None
    try:
        tms = [(int(m.group(1)), int(m.group(2) or 0), m.group(3)) for m in hits[:2]]
        if len(tms) == 2 and tms[1][2] and not tms[0][2]:
            # '2:00 - 3:30 PM' -> apply PM to start if it makes sense
            cand = _fix(tms[0][0], tms[0][1], tms[1][2])
            end = _fix(*tms[1])
            start = cand if cand < end else _fix(*tms[0])
            return start, end
        start = _fix(*tms[0])
        end = _fix(*tms[1]) if len(tms) > 1 else None
        if end and end <= start and end.hour + 12 < 24:   # '7.00-8.15' evening -> 19:00-20:15
            end = dt.time(end.hour + 12, end.minute)
        return start, end
    except ValueError:
        return None, None

def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

def load_sheets(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        grid = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:          # copy merged values into every cell
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = v
        out[ws.title] = grid
    return out

# ---------------------------------------------------------------- parsing
COLS = {
    "date": r"^\s*(date|day\s*&?\s*date|date\s*&?\s*day)\b",
    "time": r"\b(time|timing|slot)\b",
    "start": r"\b(start|from)\b",
    "end": r"\b(end|to|till)\b",
    "section": r"\b(section|sec|div|division|batch|group)\b",
    "course": r"\b(course|subject|title|module|paper)\b",
    "faculty": r"\b(faculty|instructor|professor|prof|teacher)\b",
    "room": r"\b(room|venue|classroom|class\s*room|location|hall)\b",
    "session": r"\b(session\s*(no|#|number)?|sr\.?\s*no)\b",
}

def find_header(grid):
    for i, row in enumerate(grid[:40]):
        cells = [clean(c).lower() for c in row]
        if any(re.search(COLS["date"], c) for c in cells) and sum(1 for c in cells if c) >= 2:
            return i
    return None

def map_cols(header):
    m = {}
    for j, h in enumerate(header):
        h = clean(h).lower()
        for k, pat in COLS.items():
            if k not in m and h and re.search(pat, h):
                m[k] = j
                break
    return m

def make_event(day, start, end, sec, summary, extra):
    if end is None or end <= start:
        end = (dt.datetime.combine(day, start) + dt.timedelta(minutes=DEFAULT_MINS)).time()
    return {"date": day, "start": start, "end": end, "sections": sec,
            "summary": summary or "Class", **extra}

def parse_list(grid, h, cols, year):
    events, last_date = [], None
    for row in grid[h + 1:]:
        g = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
        d = to_date(g("date"), year) or (last_date if any(clean(c) for c in row) else None)
        if not d:
            continue
        last_date = d
        sec = section_matches(clean(g("section")))
        if not sec:
            continue
        if "start" in cols:
            start, _ = parse_times(g("start")); end, _ = parse_times(g("end"))
        else:
            start, end = parse_times(g("time"))
        if not start:
            continue
        course = clean(g("course"))
        events.append(make_event(d, start, end, sec, course, {
            "faculty": clean(g("faculty")), "room": clean(g("room")), "session": clean(g("session"))}))
    return events

def find_slot_row(grid, h, date_col):
    """Time slots may be on the header row or a few rows above it; pick the row with the most."""
    best, best_n = h, -1
    for i in range(h, max(-1, h - 4), -1):
        n = sum(1 for j, c in enumerate(grid[i]) if j != date_col and parse_times(c)[0])
        if n > best_n:
            best, best_n = i, n
    return best

def bump_evening(slots):
    """Make slot start times increase left to right: '9.00-10.15' after '7.00-8.15' PM is 9 PM."""
    prev = None
    for j in sorted(slots):
        st, en = slots[j]
        if prev and st < prev and st.hour + 12 < 24:
            st = dt.time(st.hour + 12, st.minute)
            if en and en.hour + 12 < 24 and en < st:
                en = dt.time(en.hour + 12, en.minute)
        slots[j] = (st, en)
        prev = st
    return slots

CODE_RE = re.compile(r"^\s*(S\d+\s*-\s*C\d+\s*-\s*[^\s(]+)", re.I)

def split_cell(txt):
    """'S2-C1-CPB ( Darshan T) (1)' -> ('S2-C1-CPB', 'Darshan T', '1')"""
    m = CODE_RE.match(txt)
    if not m:
        return None, "", ""
    code = re.sub(r"\s+", "", m.group(1)).upper()
    rest = txt[m.end():]
    nums = re.findall(r"\d+", rest)
    fac = re.sub(r"\s+", " ", re.sub(r"[()\d]", " ", rest)).strip(" -,")
    return code, fac, (nums[-1] if nums else "")

def parse_grid(grid, h, cols, year):
    date_col = cols["date"]
    room_col = cols.get("room")
    sr = find_slot_row(grid, h, date_col)
    slots = {j: parse_times(c) for j, c in enumerate(grid[sr]) if j != date_col and j != room_col}
    slots = bump_evening({j: t for j, t in slots.items() if t[0]})
    events, last_date = [], None
    for row in grid[h + 1:]:
        d = to_date(row[date_col] if date_col < len(row) else None, year) or last_date
        if not d:
            continue
        last_date = d
        room = clean(row[room_col]) if room_col is not None and room_col < len(row) else ""
        for j, (start, end) in slots.items():
            txt = clean(row[j]) if j < len(row) else ""
            sec = section_matches(txt)
            if not sec:
                continue
            code, fac, sess = split_cell(txt)
            title = COURSE_NAMES.get(code, code or txt)
            events.append(make_event(d, start, end, sec, title, {
                "code": code or "", "faculty": fac, "session": sess, "room": room}))
    return events

COURSE_NAMES = {}

def load_course_names(sheets):
    """Map 'S2-C1-CPB' -> 'Content & Platform Business (Projects)' from any course-list sheet."""
    for grid in sheets.values():
        name_col = abbr_col = None
        for row in grid:
            cells = [clean(c).lower() for c in row]
            if "abbreviation" in cells and any("course" in c for c in cells):
                abbr_col = cells.index("abbreviation")
                name_col = next(i for i, c in enumerate(cells) if "course" in c)
                continue
            if abbr_col is not None and abbr_col < len(row):
                code = re.sub(r"\s+", "", clean(row[abbr_col])).upper()
                if CODE_RE.match(code) and clean(row[name_col]):
                    COURSE_NAMES[code] = clean(row[name_col])

def parse_workbook(sheets):
    year = dt.date.today().year
    load_course_names(sheets)
    events = []
    for name, grid in sheets.items():
        h = find_header(grid)
        if h is None:
            print(f"  sheet '{name}': no header row with a 'Date' column, skipped")
            continue
        cols = map_cols(grid[h])
        if "section" in cols:
            ev = parse_list(grid, h, cols, year); mode = "list"
        else:
            ev = parse_grid(grid, h, cols, year); mode = "grid"
        print(f"  sheet '{name}': {mode} layout, header row {h + 1}, {len(ev)} matching classes")
        events += ev
    # de-duplicate
    seen, uniq = {}, []
    for e in sorted(events, key=lambda e: (e["date"], e["start"])):
        k = (e["date"], e["start"], (e.get("code") or e["summary"]).lower())
        if k in seen:
            first = seen[k]
            if e.get("room") and e["room"] not in first.get("room", ""):
                first["room"] = " / ".join(x for x in [first.get("room"), e["room"]] if x)
            continue
        seen[k] = e; uniq.append(e)
    return uniq

# ---------------------------------------------------------------- ICS output
def esc(s):
    return str(s).replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")

def fold(line):
    out, b = [], line.encode()
    while len(b) > 75:
        cut = 75 if not out else 74
        while (b[cut] & 0xC0) == 0x80:              # don't split a UTF-8 character
            cut -= 1
        out.append(b[:cut].decode()); b = b[cut:]
    out.append(b.decode())
    return "\r\n ".join(out)

def utc(day, t):
    return dt.datetime.combine(day, t, TZ).astimezone(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")

def build_ics(events):
    L = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//mica-ics-feed//EN", "CALSCALE:GREGORIAN",
         "METHOD:PUBLISH", f"X-WR-CALNAME:{esc(CAL_NAME)}", "X-WR-TIMEZONE:" + str(TZ),
         "REFRESH-INTERVAL;VALUE=DURATION:PT1H", "X-PUBLISHED-TTL:PT1H"]
    for e in events:
        secs = "/".join(e["sections"])
        summary = f"[{secs}] {e['summary']}"
        desc = "\n".join(x for x in [
            f"Section: {secs}",
            f"Code: {e['code']}" if e.get("code") else "",
            f"Session: {e['session']}" if e.get("session") else "",
            f"Faculty: {e['faculty']}" if e.get("faculty") else "",
        ] if x)
        uid = hashlib.sha1(f"{e['date']}|{e['start']}|{(e.get('code') or e['summary']).lower()}".encode()).hexdigest()
        start = utc(e["date"], e["start"])
        L += ["BEGIN:VEVENT", f"UID:{uid}@mica-ics-feed", f"DTSTAMP:{start}",
              f"DTSTART:{start}", f"DTEND:{utc(e['date'], e['end'])}", f"SUMMARY:{esc(summary)}"]
        if e.get("room"):
            L.append(f"LOCATION:{esc(e['room'])}")
        if desc:
            L.append(f"DESCRIPTION:{esc(desc)}")
        L.append("END:VEVENT")
    L.append("END:VCALENDAR")
    return "\r\n".join(fold(l) for l in L) + "\r\n"

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", help="local .xlsx instead of downloading")
    ap.add_argument("--inspect", action="store_true", help="print first rows of each sheet")
    a = ap.parse_args()

    if a.file:
        data = open(a.file, "rb").read()
    else:
        url = os.environ.get("SHEET_URL") or sys.exit("ERROR: SHEET_URL is not set")
        data = download(url)
    sheets = load_sheets(data)

    if a.inspect:
        for name, grid in sheets.items():
            print(f"\n=== {name} ===")
            for row in grid[:12]:
                print(" | ".join(clean(c)[:25] for c in row))
        return

    print(f"Looking for sections {SECTIONS}")
    events = parse_workbook(sheets)
    print(f"Total: {len(events)} classes")

    if not events:
        if os.path.exists(OUTPUT) and "BEGIN:VEVENT" in open(OUTPUT).read():
            sys.exit("ERROR: parsed 0 classes; keeping the previous calendar instead of wiping it.")
        print("WARNING: 0 classes found. Run with --file <copy.xlsx> --inspect to check the layout.")

    os.makedirs(os.path.dirname(OUTPUT) or ".", exist_ok=True)
    with open(OUTPUT, "w", newline="") as f:
        f.write(build_ics(events))
    print(f"Wrote {OUTPUT}")

if __name__ == "__main__":
    main()
